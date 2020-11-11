# # import pdfplumber
# import pdfminer
# from pdfminer.converter import PDFPageAggregator
# from pdfminer.layout import LAParams, LTTextBoxHorizontal, LTLine, LTTextLineHorizontal, LTRect, LTCurve, LTImage, \
#     LTFigure, LTAnon, LTChar, LTTextBox, LTTextLine
# from pdfminer.pdfinterp import PDFTextExtractionNotAllowed, PDFResourceManager, PDFPageInterpreter
# from pdfminer.pdfparser import PDFParser, PDFDocument
# import pdf_tools
# import fitz

import logging
if __name__ == "__main__":
    _logger = logging.getLogger()
    logging.basicConfig(format='%(asctime)s %(module)s %(levelname)s: %(message)s',
                        datefmt='%m/%d/%Y %I:%M:%S %p')
else:
    _logger = logging.getLogger(__name__)
_logger.setLevel(logging.INFO)


file = "Compact Descriptors for Video Analysis: the Emerging MPEG Standard.pdf"
"""
q=set()

def direct(x):
    # print(type(x))
    return {
        LTLine: print,
        LTTextBoxHorizontal: print,
        LTTextLineHorizontal: print,
        LTFigure: print
    }[type(x)]

with open(file, 'rb') as pd_file:
    # 用文件对象创建一个PDF文档分析器
    parser = PDFParser(pd_file)
    # 创建一个PDF文档
    doc = PDFDocument()
    # 连接分析器，与文档对象
    parser.set_document(doc)
    doc.set_parser(parser)
    # 提供初始化密码，如果没有密码，就创建一个空的字符串
    doc.initialize()
    if not doc.is_extractable:
        raise PDFTextExtractionNotAllowed
    else:
        # 创建PDF资源管理器，来共享资源
        rsrcmgr = PDFResourceManager()
        # 创建一个PDF设备对象
        laparams = LAParams()
        device = PDFPageAggregator(rsrcmgr, laparams=laparams)
        # 创建一个PDF解释其对象
        interpreter = PDFPageInterpreter(rsrcmgr, device)
        for page in doc.get_pages():
            interpreter.process_page(page)
            # 接受该页面的LTPage对象
            layout = device.get_result()
            # 这里layout是一个LTPage对象 里面存放着 这个page解析出的各种对象
            # 一般包括LTTextBox, LTFigure, LTImage, LTTextBoxHorizontal 等等
            # 想要获取文本就获得对象的text属性，
            for x in layout:
                key = type(x)
                if key is LTLine:
                    pass
                if key is LTTextBoxHorizontal:
                    print(x.paragraphs(2))
                else:
                    print(key)
                    exit()

print(q)




import fitz
import pdfplumber
from openpyxl import Workbook
from PIL import Image
import cv2
import numpy as np
import os

def analysis_table(pdf_file):

    #打開表格
    workbook = Workbook()
    sheet = workbook.active

    #打開pdf
    with pdfplumber.open(pdf_file) as pdf:
        #遍歷每頁pdf
        for page in pdf.pages:
            #提取表格信息
            table=page.extract_table( table_settings = {
            'vertical_strategy':"text",
            "horizontal_strategy":"text"})
            print(table)

            # 格式化表格數據
            for row in table:
                print(row)
                sheet.append(row)

    workbook.save(filename="2.xlsx")

def pixmap2array(pix):
    '''pixmap數據轉數組對象'''
    #獲取顏色空間
    cspace = pix.colorspace
    if cspace is None:
        mode = "L"
    elif cspace.n == 1:
        mode = "L" if pix.alpha == 0 else "LA"
    elif cspace.n == 3:
        mode = "RGB" if pix.alpha == 0 else "RGBA"
    else:
        mode = "CMYK"

    #將byte數據轉化為PIL格式
    img = Image.frombytes(mode, (pix.width, pix.height), pix.samples)
    #將PIL轉化為numpy格式，並將RGB顏色空間轉化為BGR
    img = cv2.cvtColor(np.asarray(img), cv2.COLOR_RGB2BGR)

    return img

def ananlysis_PDF(pdf_file):
    '''解析pdf信息'''

    # 判斷pdf是否存在
    if not os.path.exists(pdf_file):
        print("pdf文件不存在")
        return;

    # 打開pdf文件
    doc=fitz.open(pdf_file)

    #遍歷pdf，提取信息
    for page in doc:
        words=page.getTextWords()
        for w in words:
            print(fitz.Rect(w[:4]),w[4])

        img_list=page.getImageList()
        i=0
        for img in img_list:
            print(fitz.Rect(img[:4]))
            pix=fitz.Pixmap(doc,img[0])
            save_name="./media/page_{}_{}.png".format(page.number,i)
            pix.writePNG(save_name)
            image=pixmap2array(pix)
            i+=1

ananlysis_PDF(file)
analysis_table(file)

"""
import fitz

doc = fitz.open(file)
from enum import Enum
import copy


class FontFlag(Enum):
    SuperScripted = 0
    Italic = 1
    Serifed = 2
    Monospaced = 3
    Bold = 4

    @classmethod
    def get_member_by_value(cls, value):
        for key, value in cls.__members__.items():
            if value.value == value:
                return value
            else:
                raise ValueError("value(%r) not in FontFlag" % value)

class WordStage(object):
    def __init__(self, bbox, font, size, flags, color, origin=None, text=None, chars=None, **kwargs):
        self.bbox = bbox
        self.origin = origin
        self.font = font
        self.size = size
        self.flags = FontFlag.get_member_by_value(flags)
        self.color = color
        self.text = text if text is not None else chars

    @staticmethod
    def _splicing(word_stages):
        first_word_stages = word_stages
        for word_stage in word_stages[:-1]:
            text = word_stage['text'].strip()
            if text[-1] == '-':
                text = text[:-1]
            first_word_stages.text.append(text)
        first_word_stages.text = ' '.join(first_word_stages.text)
        return first_word_stages

    @staticmethod
    def splicing(word_stages, unignores = ['size', 'font', 'bbox', 'color']):
        last_ws = word_stages[0]
        parts = []
        new_word_stages =[copy.deepcopy(last_ws)]
        for cur_ws in word_stages:
            for unignore in unignores:
                if getattr(last_ws, unignore, "@!#&&^*$") is not "@!#&&^*$":
                    if getattr(last_ws, unignore) != getattr(cur_ws, unignore):
                        parts.append([copy.deepcopy(cur_ws)])
                        new_word_stages = [copy.deepcopy(cur_ws)]
                        last_ws = cur_ws
                    else:
                        new_word_stages.append(copy.deepcopy(cur_ws))

        first_word_stages = word_stages[0]
        first_word_stages.text = [first_word_stages.text, ]
        for word_stage in word_stages[1:]:
            if word_stage.size != first_word_stages.size:
                _logger.debug("first_word_stages.size(%r) != word_stage.size(%r)" % (word_stage.size, first_word_stages.size))

            if word_stage.size != first_word_stages.size:
                _logger.debug("first_word_stages.size(%r) != word_stage.size(%r)" % (word_stage.size, first_word_stages.size))

            if word_stage.size != first_word_stages.size:
                _logger.debug("first_word_stages.size(%r) != word_stage.size(%r)" % (word_stage.size, first_word_stages.size))

            if word_stage.size != first_word_stages.size:
                _logger.debug("first_word_stages.size(%r) != word_stage.size(%r)" % (word_stage.size, first_word_stages.size))

            text = word_stage['text'].strip()
            if text[-1] == '-':
                text = text[:-1]
            first_word_stages.text.append(text)
        first_word_stages.text = ' '.join(first_word_stages.text)
        return first_word_stages







class PDFReader(object):
    def close(self):
        if isinstance(self.doc, fitz.Document) and not self.doc.isClosed:
            self.doc.close()

    @staticmethod
    def open(filepath):
        pdf_doc = fitz.open(filepath)
        if pdf_doc.isPDF:
            return pdf_doc
        raise

    @property
    def filepath(self):
        return self._filepath

    @filepath.setter
    def filepath(self, filepath):
        self.close()
        pdf_doc = self.open(filepath)
        self.close()
        self.doc = pdf_doc
        self._filepath = filepath

    def __delete__(self):
        self.close()

    def __enter__(self, filepath):
        return self

    def __exit__(self, exc_type, exc_val, exc_tb):
        self.close()

    def __init__(self, filepath):
        self.filepath = filepath
        self.doc = fitz.open(filepath)

    def get_contents(self):
        return self.doc.getToC()[0]

    def get_page_count(self):
        return self.doc.pageCount

    def parse(self):
        def parse_image_block(block):
            # block['type']==1
            width = block['width']
            height = block['height']
            number = block['number'] # 索引
            ext = block['ext']
            colorspace = block['colorspace'] # 色彩空间
            image = block['image'] # b''

def parse_text_lines(lines):
    # block['type']==0
    part = []
    for line in lines:
        text_line = []
        span_font = line['spans'][0]['font']
        for span in line['spans']:
            if span_font != span['font']:
                _logger.debug("span_font: {}, span['font']: {}".format(span_font, span['font']))
            text = span['text'].strip()
            if text[-1]=='-':
                text = text[:-1] + ' '
            else:
                text = text + ' '
            text_line.append(text)
        text_line = ''.join(text_line)
        part.append(text_line)
    print(part)
    # part = ''.join(part)


        # for page in self.doc.pages(start, stop, step):
        data = []
        for page in self.doc.pages():
            page_dict = page.getText('dict')
            if 'blocks' in page_dict:
                for block in page_dict['blocks']:



